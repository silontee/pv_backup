# -*- coding: utf-8 -*-
"""
PRD v4 PPT — 흰 배경 + 초록 테마 스타일리시 버전
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# ── 색상 팔레트 (초록 기반) ──
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE = RGBColor(0xFA, 0xFA, 0xFA)
LIGHT_BG = RGBColor(0xF5, 0xF7, 0xF5)       # 아주 연한 초록 회색
GREEN_DARK = RGBColor(0x0B, 0x3D, 0x2E)      # 짙은 초록 (텍스트/타이틀)
GREEN_PRIMARY = RGBColor(0x00, 0x7A, 0x5E)   # 메인 초록
GREEN_ACCENT = RGBColor(0x00, 0xA6, 0x7E)    # 밝은 초록
GREEN_LIGHT = RGBColor(0xD0, 0xEC, 0xE2)     # 연한 초록 (카드 배경)
GREEN_PALE = RGBColor(0xE8, 0xF5, 0xEF)      # 더 연한 초록
MINT = RGBColor(0x4E, 0xD4, 0xA1)            # 민트
DARK_TEXT = RGBColor(0x2C, 0x2C, 0x2C)       # 본문 텍스트
MID_TEXT = RGBColor(0x5A, 0x5A, 0x5A)        # 서브 텍스트
LIGHT_TEXT = RGBColor(0x8A, 0x8A, 0x8A)      # 약한 텍스트
ORANGE_WARN = RGBColor(0xE8, 0x8D, 0x2A)     # 경고 주황
RED_ALERT = RGBColor(0xDC, 0x3A, 0x3A)       # 위험 빨강
BLUE_INFO = RGBColor(0x38, 0x8E, 0xD1)       # 정보 파랑
CHARCOAL = RGBColor(0x1A, 0x1A, 0x1A)        # 거의 검정

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_bg(slide, color=WHITE):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def add_rounded(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def add_circle(slide, left, top, size, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def tb(slide, left, top, width, height, text, size=18, color=DARK_TEXT, bold=False, align=PP_ALIGN.LEFT, font="맑은 고딕"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font
    p.alignment = align
    return txBox


def bullets(slide, left, top, width, height, items, size=14, color=MID_TEXT, spacing=Pt(6)):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.name = "맑은 고딕"
        p.space_after = spacing
    return txBox


def card(slide, left, top, width, height, title, body_lines, accent=GREEN_PRIMARY, bg_color=GREEN_PALE):
    add_rounded(slide, left, top, width, height, bg_color)
    add_rect(slide, left, top, width, Inches(0.05), accent)
    tb(slide, left + Inches(0.3), top + Inches(0.2), width - Inches(0.6), Inches(0.4),
       title, size=16, color=accent, bold=True)
    y = top + Inches(0.6)
    for line in body_lines:
        if line == "":
            y += Inches(0.1)
            continue
        tb(slide, left + Inches(0.3), y, width - Inches(0.6), Inches(0.3),
           line, size=12, color=DARK_TEXT)
        y += Inches(0.27)


def section_header(slide, number, title):
    """슬라이드 상단 섹션 헤더"""
    # 상단 초록 바
    add_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), GREEN_PRIMARY)
    # 번호 원
    c = add_circle(slide, Inches(0.8), Inches(0.5), Inches(0.55), GREEN_PRIMARY)
    tb(slide, Inches(0.8), Inches(0.55), Inches(0.55), Inches(0.45),
       number, size=22, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    # 제목
    tb(slide, Inches(1.55), Inches(0.5), Inches(10), Inches(0.6),
       title, size=30, color=GREEN_DARK, bold=True)
    # 구분선
    add_rect(slide, Inches(0.8), Inches(1.2), Inches(3), Inches(0.03), GREEN_ACCENT)


# ============================================================
# SLIDE 1: 표지
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)

# 좌측 초록 패널
add_rect(slide, Inches(0), Inches(0), Inches(5.5), Inches(7.5), GREEN_DARK)

# 좌측 장식 원
add_circle(slide, Inches(-0.5), Inches(5.5), Inches(3), RGBColor(0x0D, 0x4A, 0x38))
add_circle(slide, Inches(3.5), Inches(-0.5), Inches(2), RGBColor(0x0D, 0x4A, 0x38))

# 좌측 텍스트
tb(slide, Inches(0.8), Inches(2.0), Inches(4), Inches(0.5),
   "한국남동발전", size=18, color=MINT)
tb(slide, Inches(0.8), Inches(2.5), Inches(4.2), Inches(2.0),
   "태양광 변동성 대비\n화력 백업 최적화\n플랫폼", size=38, color=WHITE, bold=True)
tb(slide, Inches(0.8), Inches(4.8), Inches(4), Inches(0.5),
   "내부 의사결정 지원 시스템", size=16, color=RGBColor(0x7C, 0xC4, 0xA8))

# 우측 내용
tb(slide, Inches(6.5), Inches(2.5), Inches(6), Inches(0.5),
   "Product Requirements Document", size=20, color=GREEN_PRIMARY, bold=True)

features = [
    "태양광 변동성 정밀 예측",
    "화력 포트폴리오 최적 대응 계획",
    "사전 대비 → 실시간 감시 → 실시간 대응",
    "ESG 대기질 영향 시각화",
]
y = Inches(3.3)
for feat in features:
    add_circle(slide, Inches(6.5), y + Inches(0.05), Inches(0.15), GREEN_ACCENT)
    tb(slide, Inches(6.9), y, Inches(5.5), Inches(0.4),
       feat, size=16, color=DARK_TEXT)
    y += Inches(0.5)

# 버전 정보
add_rect(slide, Inches(6.5), Inches(6.0), Inches(2.5), Inches(0.5), GREEN_PALE)
tb(slide, Inches(6.5), Inches(6.05), Inches(2.5), Inches(0.4),
   "  PRD v4  ·  2026.04", size=14, color=GREEN_PRIMARY, bold=True)

# ============================================================
# SLIDE 2: 왜 필요한가
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
section_header(slide, "01", "왜 이 플랫폼이 필요한가")

# 3개 카드
card(slide, Inches(0.8), Inches(1.6), Inches(3.6), Inches(2.7),
     "☀  태양광: 무조건 최대 발전", [
         "변동비 0원 → 최우선 급전",
         "나오는 만큼 전량 입찰 (원칙)",
         "",
         "BUT: 내일 얼마 나올지 모른다",
         "구름 많은 날 오차 15~30%",
         "급변 날씨 오차 30~60%",
     ], accent=ORANGE_WARN, bg_color=RGBColor(0xFF, 0xF8, 0xEB))

card(slide, Inches(4.8), Inches(1.6), Inches(3.6), Inches(2.7),
     "🏭  화력: 이미 돌아가고 있다", [
         "최소출력 합계 ~6,387MW가 24시간 가동",
         "석탄 60%, LNG 48% 이하로 못 내림",
         "",
         "출력 정밀도: ±1~3% (매우 정확)",
         "BUT: 미리 준비해야 함",
         "석탄 Cold 기동: 12~24시간",
     ], accent=GREEN_PRIMARY, bg_color=GREEN_PALE)

card(slide, Inches(8.8), Inches(1.6), Inches(3.6), Inches(2.7),
     "⚡  충돌: 불확실 × 느림", [
         "태양광 급감 → 화력 올려야",
         "  석탄 램프: 9~26MW/분",
         "",
         "태양광 과잉 → 화력 내려야",
         "  최소출력 이하로 못 내림",
         "  끄면 재기동 수 시간",
     ], accent=RED_ALERT, bg_color=RGBColor(0xFD, 0xF0, 0xF0))

# 하단 핵심 메시지
add_rounded(slide, Inches(0.8), Inches(4.7), Inches(11.7), Inches(2.3), GREEN_DARK)
tb(slide, Inches(1.5), Inches(4.9), Inches(10.3), Inches(0.4),
   "핵심 인사이트", size=18, color=MINT, bold=True)
tb(slide, Inches(1.5), Inches(5.35), Inches(10.3), Inches(1.5),
   "태양광은 \"내일 얼마 나올지 모르겠다\"  →  불확실\n"
   "화력은 \"얼마를 내라고 하면 정확히 내지만, 미리 말해줘야 한다\"  →  정확하지만 느림\n\n"
   "태양광 예측을 잘 해서, 화력에게 미리 정확한 지시를 내리는 것이 핵심\n"
   "이 지시 = D-1일 가용용량 신고 = 입찰 최적화",
   size=15, color=RGBColor(0xCC, 0xE5, 0xD9))

# ============================================================
# SLIDE 3: 현행 제도 & 벤치마크
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
section_header(slide, "02", "현행 제도와 경쟁 환경")

# CBP
card(slide, Inches(0.8), Inches(1.6), Inches(5.5), Inches(2.3),
     "한국 전력시장 (CBP) 입찰 프로세스", [
         "D-1일 10:00   1차 예측 & 가용용량 신고",
         "D-1일 17:00   2차 예측 (최신 기상으로 보정)",
         "D-1일 저녁     KPX 급전계획 확정 → 발전사 통보",
         "D일 실시간     급전지시에 따라 발전 수행",
         "",
         "변동비 기반 → 가격이 아닌 가용용량 최적화가 핵심",
     ], accent=GREEN_PRIMARY, bg_color=GREEN_PALE)

card(slide, Inches(6.8), Inches(1.6), Inches(5.5), Inches(2.3),
     "재생에너지 예측제도 (2021.10~)", [
         "20MW 초과 태양광/풍력 → 예측 제출 의무",
         "D-1일 10시, 17시 2차에 걸쳐 제출",
         "기준오차율 충족 시 인센티브 / 미충족 시 페널티",
         "",
         "→ 남동발전도 이미 예측하고 있음",
         "→ 하지만 구체적 시스템/정확도 미공개",
     ], accent=BLUE_INFO, bg_color=RGBColor(0xEA, 0xF4, 0xFB))

# 벤치마크 테이블
add_rounded(slide, Inches(0.8), Inches(4.3), Inches(11.7), Inches(2.6), LIGHT_BG)
tb(slide, Inches(1.2), Inches(4.4), Inches(4), Inches(0.4),
   "경쟁 벤치마크", size=18, color=GREEN_DARK, bold=True)

bench_h = ["발전사", "예측 시스템", "정확도", "화력 연계"]
bench_d = [
    ["한국중부발전", "AI (제이케이코어 공동)", "99% (2026.2 발표)", "2030 통합 운영 목표"],
    ["한전 (KEPCO)", "AI 재생에너지 예측 기술", "2km 격자, 1주일 예측", "개발 완료 (2021)"],
    ["한국남동발전", "미공개", "미공개", "본 플랫폼이 해결 →"],
]

y = Inches(4.85)
for ci, h in enumerate(bench_h):
    x = Inches(1.2 + ci * 2.7)
    tb(slide, x, y, Inches(2.5), Inches(0.3), h, size=13, color=GREEN_PRIMARY, bold=True)

for ri, row in enumerate(bench_d):
    bg_c = WHITE if ri % 2 == 0 else LIGHT_BG
    add_rect(slide, Inches(1.0), y + Inches(0.35 + ri * 0.5), Inches(11.2), Inches(0.45), bg_c)
    for ci, cell in enumerate(row):
        x = Inches(1.2 + ci * 2.7)
        yt = y + Inches(0.38 + ri * 0.5)
        c = GREEN_PRIMARY if ri == 2 and ci >= 2 else DARK_TEXT
        b = True if ri == 2 and ci >= 2 else False
        tb(slide, x, yt, Inches(2.5), Inches(0.35), cell, size=12, color=c, bold=b)

# ============================================================
# SLIDE 4: 3단계 운영 모드
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
section_header(slide, "03", "플랫폼 3단계 운영 모드")

steps = [
    ("STEP 1", "사전 대비", "D-1일", GREEN_PRIMARY, GREEN_PALE, [
        "D-1 10:00  1차 예측 & 대응 계획",
        "D-1 17:00  2차 보정",
        "",
        "기상 데이터:",
        "  LDAPS 수치예보 (하루 8회)",
        "  단기예보 (하루 8회)",
        "  초단기예보 (30분마다)",
        "",
        "산출물:",
        "  24시간 태양광 예측",
        "  화력 액션 플랜 타임라인",
        "  시나리오별 비교표",
        "  KPX 가용용량 신고안",
    ]),
    ("STEP 2", "실시간 감시", "D일", BLUE_INFO, RGBColor(0xEA, 0xF4, 0xFB), [
        "매 30분 예측 갱신",
        "매 10분 위성 구름 확인",
        "",
        "데이터:",
        "  초단기예보 (30분마다)",
        "  위성 영상 (10분마다)",
        "  발전소 기상 실측 (실시간)",
        "  태양광 발전 실측 (실시간)",
        "",
        "핵심:",
        "  예측 vs 실측 괴리 감지",
        "  향후 3~6시간 예측 갱신",
    ]),
    ("STEP 3", "실시간 대응", "이벤트 발생 시", RED_ALERT, RGBColor(0xFD, 0xF0, 0xF0), [
        '"태양광이 예측보다',
        ' 300MW 적게 나오고 있습니다"',
        "",
        "시스템 자동:",
        "  최신 기상으로 재계산",
        "  화력 조정 지시 즉시 갱신",
        "",
        "출력 예시:",
        '  "영흥 2호 지금 출력 증가"',
        '  "분당 1호 기동 준비 (45분)"',
        "",
        "→ 사후 대비가 핵심 가치",
    ]),
]

for i, (step_no, step_name, timing, accent, bg_c, lines) in enumerate(steps):
    x = Inches(0.6 + i * 4.1)
    # 카드
    add_rounded(slide, x, Inches(1.5), Inches(3.8), Inches(5.2), bg_c)
    add_rect(slide, x, Inches(1.5), Inches(3.8), Inches(0.06), accent)
    # 스텝 배지
    add_rounded(slide, x + Inches(0.2), Inches(1.7), Inches(1.2), Inches(0.4), accent)
    tb(slide, x + Inches(0.2), Inches(1.72), Inches(1.2), Inches(0.35),
       step_no, size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    # 제목
    tb(slide, x + Inches(1.5), Inches(1.72), Inches(2.0), Inches(0.35),
       step_name, size=18, color=accent, bold=True)
    # 타이밍
    tb(slide, x + Inches(0.3), Inches(2.2), Inches(3.2), Inches(0.3),
       timing, size=12, color=LIGHT_TEXT)
    # 본문
    y = Inches(2.6)
    for line in lines:
        if line == "":
            y += Inches(0.08)
            continue
        tb(slide, x + Inches(0.3), y, Inches(3.2), Inches(0.25),
           line, size=11, color=DARK_TEXT)
        y += Inches(0.25)

    # 화살표
    if i < 2:
        tb(slide, x + Inches(3.8), Inches(3.5), Inches(0.3), Inches(0.5),
           "→", size=28, color=GREEN_ACCENT, align=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 5: 기상 데이터 전략
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
section_header(slide, "04", "기상 데이터 전략: 기상청 오차 극복")

# 문제 제기
add_rounded(slide, Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.7), RGBColor(0xFF, 0xF8, 0xEB))
tb(slide, Inches(1.2), Inches(1.55), Inches(11), Inches(0.6),
   '"기상청이 틀린만큼 틀리는 것 아닌가?"  →  기상청만 쓰면 그렇다. 다중 소스를 쓰면 극복 가능.',
   size=15, color=ORANGE_WARN, bold=True)

# 데이터 계층 테이블
layers = [
    ("LDAPS 수치예보", "1.5km", "하루 8회", "48시간", "하루전 예측 주입력", GREEN_PRIMARY),
    ("단기예보 (동네예보)", "5km", "하루 8회", "3일", "하루전 예측 보조", GREEN_PRIMARY),
    ("초단기예보", "5km", "30분마다", "6시간", "실시간 보정 핵심", GREEN_ACCENT),
    ("위성 영상 (천리안2A)", "2km", "10분마다", "실시간", "구름 급변 탐지", MINT),
    ("발전소 현장 관측", "포인트", "실시간", "현재", "실측 기반 보정", MINT),
    ("태양광 발전 실측", "인버터별", "실시간", "현재", "예측 vs 실측 비교", MINT),
]

headers_met = ["데이터 소스", "해상도", "갱신 주기", "범위", "용도"]
y = Inches(2.5)
for ci, h in enumerate(headers_met):
    x = Inches(0.8 + ci * 2.4)
    tb(slide, x, y, Inches(2.2), Inches(0.35), h, size=13, color=GREEN_PRIMARY, bold=True)

for ri, (name, res, freq, rng, use, clr) in enumerate(layers):
    bg_c = GREEN_PALE if ri % 2 == 0 else WHITE
    add_rect(slide, Inches(0.7), y + Inches(0.4 + ri * 0.48), Inches(12), Inches(0.44), bg_c)
    row = [name, res, freq, rng, use]
    for ci, cell in enumerate(row):
        x = Inches(0.8 + ci * 2.4)
        yt = y + Inches(0.42 + ri * 0.48)
        c = clr if ci == 2 and ("실시간" in cell or "분마다" in cell) else DARK_TEXT
        b = True if ci == 2 and ("실시간" in cell or "분마다" in cell) else False
        tb(slide, x, yt, Inches(2.2), Inches(0.35), cell, size=12, color=c, bold=b)

# 극복 방법
add_rounded(slide, Inches(0.8), Inches(5.6), Inches(11.7), Inches(1.5), GREEN_DARK)
tb(slide, Inches(1.3), Inches(5.7), Inches(10.7), Inches(0.35),
   "극복 전략", size=16, color=MINT, bold=True)
tb(slide, Inches(1.3), Inches(6.05), Inches(10.7), Inches(0.95),
   "① 다중 소스 앙상블: LDAPS + 위성 + 실측 결합 → 단일 소스 의존도 감소\n"
   "② 위성 기반 초단기: 구름 이동 직접 관측 → 2~3시간 예측에서 NWP보다 우수\n"
   "③ 실측 피드백: 현재 예측-실측 차이를 실시간 보정\n"
   "④ 과거 패턴 학습: \"이 기상 조건에서 기상청이 얼마나 틀렸는지\" 학습",
   size=13, color=RGBColor(0xCC, 0xE5, 0xD9))

# ============================================================
# SLIDE 6: 태양광 예측 모델
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
section_header(slide, "05", "태양광 예측 모델")

# 모델 비교
models = [
    ("통계\n(회귀/ARIMA)", "15~25%", LIGHT_TEXT),
    ("ML\n(XGBoost)", "10~18%", ORANGE_WARN),
    ("DL\n(LSTM/GRU)", "8~15%", BLUE_INFO),
    ("베이지안\n+GRU (제안)", "7~12%", GREEN_PRIMARY),
    ("위성\nNowcast", "5~10%", MINT),
]

for i, (name, acc, color) in enumerate(models):
    x = Inches(0.6 + i * 2.45)
    # 카드
    bg = GREEN_PALE if i == 3 else LIGHT_BG
    add_rounded(slide, x, Inches(1.5), Inches(2.2), Inches(2.0), bg)
    if i == 3:
        add_rect(slide, x, Inches(1.5), Inches(2.2), Inches(0.06), GREEN_PRIMARY)
    tb(slide, x + Inches(0.1), Inches(1.65), Inches(2.0), Inches(0.7),
       name, size=14, color=DARK_TEXT, bold=True, align=PP_ALIGN.CENTER)
    tb(slide, x + Inches(0.1), Inches(2.4), Inches(2.0), Inches(0.3),
       "nRMSE", size=11, color=LIGHT_TEXT, align=PP_ALIGN.CENTER)
    tb(slide, x + Inches(0.1), Inches(2.65), Inches(2.0), Inches(0.5),
       acc, size=22, color=color, bold=True, align=PP_ALIGN.CENTER)

# MVP 전략
add_rounded(slide, Inches(0.8), Inches(3.8), Inches(11.7), Inches(0.8), RGBColor(0xFF, 0xF8, 0xEB))
tb(slide, Inches(1.2), Inches(3.9), Inches(10.7), Inches(0.6),
   "MVP 전략:  검증된 ML(XGBoost/LSTM)로 시작  →  데이터 축적 후 베이지안+GRU 확장\n"
   "전제 조건:  남동발전 기존 예측보다 정확하지 않으면 가치 없음",
   size=14, color=DARK_TEXT)

# 2단계 모델
card(slide, Inches(0.8), Inches(4.9), Inches(5.5), Inches(2.1),
     "1단계: 베이지안 계층모형", [
         "지역별 기상 민감도 + 시간/계절 구조 posterior 추정",
         "부분풀링 → 데이터 적은 사이트도 안정 추정",
         "베타 회귀 + Fourier basis expansion",
         "출력: baseline 예측 + 잔차",
     ], accent=GREEN_PRIMARY, bg_color=GREEN_PALE)

card(slide, Inches(6.8), Inches(4.9), Inches(5.5), Inches(2.1),
     "2단계: FiLM-조건부 GRU", [
         "구름 급변, NWP 오차, ramp event 비선형 패턴 학습",
         "1단계 posterior 요약 → FiLM 조건벡터 주입",
         "같은 backbone이 지역별로 다르게 동작",
         "Direct multi-horizon 출력 (1~24시간)",
     ], accent=BLUE_INFO, bg_color=RGBColor(0xEA, 0xF4, 0xFB))

# ============================================================
# SLIDE 7: 화력 대응 계획 (핵심)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
section_header(slide, "06", "화력 대응 계획 생성 — 핵심 기능")

tb(slide, Inches(0.8), Inches(1.3), Inches(11), Inches(0.4),
   "태양광 예측이 입력  →  나오는 결과는 \"발전기별로 몇 시에 무슨 행동을 하라\"는 액션 플랜",
   size=15, color=GREEN_PRIMARY, bold=True)

# 흐름도
flow = [
    ("태양광\n예측", "내일 시간대별\n발전량", ORANGE_WARN),
    ("갭\n계산", "수요-태양광\n-고정화력", BLUE_INFO),
    ("UC\n최적화", "MILP 24시간\n최적 조합", GREEN_PRIMARY),
    ("역산\n타이밍", "발전기별\n\"몇시에 몇MW\"", GREEN_ACCENT),
    ("액션\n플랜", "운영자용\n타임라인", GREEN_DARK),
]

for i, (title, desc, color) in enumerate(flow):
    x = Inches(0.5 + i * 2.5)
    add_rounded(slide, x, Inches(1.8), Inches(2.2), Inches(1.5), GREEN_PALE)
    add_rect(slide, x, Inches(1.8), Inches(2.2), Inches(0.05), color)
    tb(slide, x + Inches(0.1), Inches(1.9), Inches(2.0), Inches(0.6),
       title, size=15, color=color, bold=True, align=PP_ALIGN.CENTER)
    tb(slide, x + Inches(0.1), Inches(2.55), Inches(2.0), Inches(0.6),
       desc, size=11, color=MID_TEXT, align=PP_ALIGN.CENTER)
    if i < 4:
        tb(slide, x + Inches(2.15), Inches(2.2), Inches(0.4), Inches(0.5),
           "→", size=22, color=GREEN_ACCENT, bold=True, align=PP_ALIGN.CENTER)

# 액션 플랜 테이블
add_rounded(slide, Inches(0.8), Inches(3.6), Inches(11.7), Inches(3.5), LIGHT_BG)
tb(slide, Inches(1.2), Inches(3.7), Inches(4), Inches(0.35),
   "액션 플랜 출력 예시", size=16, color=GREEN_DARK, bold=True)

tl_headers = ["시간", "발전기", "액션", "목표 출력"]
tl_data = [
    ("12:37", "영흥 1호", "출력 증가 시작", "520 → 700MW"),
    ("13:00", "영흥 1호", "700MW 도달", "700MW"),
    ("13:30", "영흥 4호", "출력 증가 시작", "520 → 870MW"),
    ("14:00", "분당 1호", "기동 준비 시작", "(정지 → 기동)"),
    ("14:23", "영흥 4호", "870MW 도달", "870MW (최대)"),
    ("14:45", "분당 1호", "최소출력 도달", "221MW"),
    ("15:06", "분당 1호", "400MW 도달", "400MW"),
]

y_base = Inches(4.1)
for ci, h in enumerate(tl_headers):
    x = Inches(1.2 + ci * 2.7)
    tb(slide, x, y_base, Inches(2.5), Inches(0.3), h, size=12, color=GREEN_PRIMARY, bold=True)

for ri, row in enumerate(tl_data):
    bg_c = WHITE if ri % 2 == 0 else LIGHT_BG
    add_rect(slide, Inches(1.0), y_base + Inches(0.35 + ri * 0.38), Inches(11.2), Inches(0.35), bg_c)
    for ci, cell in enumerate(row):
        x = Inches(1.2 + ci * 2.7)
        yt = y_base + Inches(0.37 + ri * 0.38)
        c = GREEN_ACCENT if "도달" in cell else ORANGE_WARN if "시작" in cell else DARK_TEXT
        tb(slide, x, yt, Inches(2.5), Inches(0.3), cell, size=11, color=c)

# 요약 바
add_rounded(slide, Inches(1.0), Inches(6.85), Inches(11.2), Inches(0.35), GREEN_PALE)
tb(slide, Inches(1.3), Inches(6.87), Inches(10.5), Inches(0.3),
   "추가 연료비: ~X억원  |  예비력 최소: 15시 350MW  |  ESG: 영흥 NOx +12%",
   size=12, color=GREEN_PRIMARY, bold=True, align=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 8: 시나리오 비교
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
section_header(slide, "07", "시나리오 비교 — 운영자 의사결정")

scen_h = ["시나리오", "태양광 가정", "추가 기동", "연료비", "예비력", "대응 실패"]
scen_d = [
    ("낙관", "예측 +10%", "없음", "기본", "충분", "없음", GREEN_ACCENT, GREEN_PALE),
    ("중립", "예측 그대로", "영흥1호 증가", "+2.1억", "충분", "없음", BLUE_INFO, RGBColor(0xEA, 0xF4, 0xFB)),
    ("비관", "예측 -30%", "영흥1,4호+분당", "+5.8억", "주의", "없음", ORANGE_WARN, RGBColor(0xFF, 0xF8, 0xEB)),
    ("급변", "14시 50% 감소", "분당 즉시 기동", "+7.2억", "위험", "지연 가능", RED_ALERT, RGBColor(0xFD, 0xF0, 0xF0)),
]

y = Inches(1.6)
for ci, h in enumerate(scen_h):
    x = Inches(0.8 + ci * 2.0)
    tb(slide, x, y, Inches(1.9), Inches(0.35), h, size=14, color=GREEN_DARK, bold=True)

for ri, (name, solar, gen, cost, reserve, fail, accent, bg_c) in enumerate(scen_d):
    add_rounded(slide, Inches(0.6), y + Inches(0.45 + ri * 0.7), Inches(12.1), Inches(0.6), bg_c)
    add_rect(slide, Inches(0.6), y + Inches(0.45 + ri * 0.7), Inches(0.06), Inches(0.6), accent)
    row = [name, solar, gen, cost, reserve, fail]
    for ci, cell in enumerate(row):
        x = Inches(0.8 + ci * 2.0)
        yt = y + Inches(0.52 + ri * 0.7)
        c = accent if ci == 0 else (RED_ALERT if "위험" in cell or "지연" in cell else DARK_TEXT)
        b = ci == 0
        tb(slide, x, yt, Inches(1.9), Inches(0.35), cell, size=13, color=c, bold=b)

# 의사결정 포인트
add_rounded(slide, Inches(0.8), Inches(4.8), Inches(11.7), Inches(2.2), GREEN_DARK)
tb(slide, Inches(1.3), Inches(4.95), Inches(10.7), Inches(0.35),
   "운영자 의사결정 포인트", size=18, color=MINT, bold=True)
tb(slide, Inches(1.3), Inches(5.35), Inches(10.7), Inches(1.5),
   "\"비관 시나리오 대비 시 연료비 3.7억 더 들지만, 급전 대응 실패 리스크 0.\n"
   " 오늘 날씨가 불확실하니 비관으로 간다.\"\n\n"
   "시스템은 추천하고, 사람이 최종 판단한다.\n"
   "비용 / 예비력 / ESG 영향을 한 화면에서 비교할 수 있게 하는 것이 대시보드의 핵심 가치.",
   size=15, color=RGBColor(0xCC, 0xE5, 0xD9))

# ============================================================
# SLIDE 9: 화력 운전 제약
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
section_header(slide, "08", "화력 발전기 운전 제약")

p_headers = ["파라미터", "영흥 (석탄)", "삼천포 (석탄)", "분당 (LNG)", "의미"]
p_data = [
    ["정격출력", "870MW/호기", "560MW/호기", "460MW", "최대 출력"],
    ["최소출력", "520MW (60%)", "336MW (60%)", "221MW (48%)", "이하로 못 내림"],
    ["램프업", "~15MW/분", "~10MW/분", "~30MW/분", "분당 증가 한계"],
    ["최소기동시간", "12시간", "12시간", "4시간", "켜면 최소 유지"],
    ["최소정지시간", "8시간", "8시간", "2시간", "끄면 최소 유지"],
    ["Cold 기동", "16시간", "16시간", "6시간", "완전 정지 후 재가동"],
    ["기동비용", "~3억원/회", "~2억원/회", "~5천만원/회", "1회 기동 비용"],
    ["변동비", "~80원/kWh", "~80원/kWh", "~160원/kWh", "연료비 기반"],
]

y = Inches(1.5)
for ci, h in enumerate(p_headers):
    x = Inches(0.6 + ci * 2.5)
    tb(slide, x, y, Inches(2.3), Inches(0.35), h, size=13, color=GREEN_PRIMARY, bold=True)

for ri, row in enumerate(p_data):
    bg_c = GREEN_PALE if ri % 2 == 0 else WHITE
    add_rect(slide, Inches(0.5), y + Inches(0.4 + ri * 0.45), Inches(12.3), Inches(0.42), bg_c)
    for ci, cell in enumerate(row):
        x = Inches(0.6 + ci * 2.5)
        yt = y + Inches(0.42 + ri * 0.45)
        c = ORANGE_WARN if ci == 3 and "160" in cell else DARK_TEXT
        tb(slide, x, yt, Inches(2.3), Inches(0.35), cell, size=12, color=c)

# trade-off 박스
add_rounded(slide, Inches(0.8), Inches(5.5), Inches(11.7), Inches(1.5), GREEN_DARK)
tb(slide, Inches(1.3), Inches(5.65), Inches(10.7), Inches(0.35),
   "핵심 Trade-off", size=16, color=MINT, bold=True)
tb(slide, Inches(1.3), Inches(6.0), Inches(10.7), Inches(0.8),
   "석탄 = 싸지만 느리다  (램프 15MW/분, Cold 기동 16시간)\n"
   "LNG = 비싸지만 빠르다  (램프 30MW/분, Cold 기동 6시간, 변동비 2배)\n\n"
   "→ 태양광 급변 대비에 LNG가 유리하지만 비용 2배. 이 균형을 최적화하는 것이 UC의 역할.",
   size=14, color=RGBColor(0xCC, 0xE5, 0xD9))

# ============================================================
# SLIDE 10: 포트폴리오 3영역
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
section_header(slide, "09", "남동발전 발전 포트폴리오 구조")

card(slide, Inches(0.8), Inches(1.5), Inches(3.6), Inches(3.5),
     "고정 영역", [
         "건드릴 수 없음",
         "",
         "필수운전 + 최소출력 구간",
         "합계: ~6,387MW",
         "",
         "영흥: 최소 3,048MW",
         "삼천포: 최소 1,248MW",
         "강릉: 최소 1,248MW",
         "분당: 최소 442MW",
         "여수: 최소 401MW",
     ], accent=LIGHT_TEXT, bg_color=LIGHT_BG)

card(slide, Inches(4.8), Inches(1.5), Inches(3.6), Inches(3.5),
     "예측 영역 (변동)", [
         "예측이 핵심 가치",
         "",
         "태양광: 20+ 사이트",
         "  → 변동성 큼, 예측 대상",
         "",
         "풍력: 영흥, 군위 등",
         "  → 변동성 큼",
         "",
         "연료전지: 비교적 안정",
         "해양소수력: 소규모, 안정",
     ], accent=ORANGE_WARN, bg_color=RGBColor(0xFF, 0xF8, 0xEB))

card(slide, Inches(8.8), Inches(1.5), Inches(3.6), Inches(3.5),
     "최적화 영역 (조정 가능)", [
         "이 영역이 최적화 대상",
         "",
         "최소출력 ~ 최대출력 사이",
         "조정 가능: ~4,442MW",
         "",
         "영흥: +2,032MW 여유",
         "삼천포: +832MW 여유",
         "분당: +478MW 여유",
         "",
         "대기 발전기 기동/정지 결정",
     ], accent=GREEN_PRIMARY, bg_color=GREEN_PALE)

# 수식
add_rounded(slide, Inches(0.8), Inches(5.3), Inches(11.7), Inches(1.7), GREEN_DARK)
tb(slide, Inches(1.3), Inches(5.5), Inches(10.7), Inches(1.3),
   "전체 공급  =  고정 영역 (~6,387MW)  +  태양광 예측 (변동)  +  최적화 영역 (0 ~ 4,442MW)\n\n"
   "태양광이 많이 나오면 → 최적화 영역을 줄임 (연료비 절감)\n"
   "태양광이 적게 나오면 → 최적화 영역을 올림 (수요 충족)\n"
   "태양광이 예측보다 적으면 → 실시간으로 최적화 영역을 올림 (사후 대응)",
   size=14, color=WHITE)

# ============================================================
# SLIDE 11: ESG
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
section_header(slide, "10", "ESG: 하이퍼로컬 대기질 영향")

card(slide, Inches(0.8), Inches(1.5), Inches(3.6), Inches(2.8),
     "배출량 추정", [
         "",
         "화력 출력 × 배출계수",
         "= SOx, NOx, PM 배출량",
         "",
         "데이터:",
         "  남동발전 대기오염물질 배출실적",
         "  발전기별 배출계수",
     ], accent=ORANGE_WARN, bg_color=RGBColor(0xFF, 0xF8, 0xEB))

card(slide, Inches(4.8), Inches(1.5), Inches(3.6), Inches(2.8),
     "확산 모델", [
         "",
         "풍속, 풍향, 대기안정도",
         "→ 가우시안 확산 모델",
         "→ 시간대별 영향권 계산",
         "",
         "데이터:",
         "  발전소 기상정보 (풍향/풍속)",
     ], accent=BLUE_INFO, bg_color=RGBColor(0xEA, 0xF4, 0xFB))

card(slide, Inches(8.8), Inches(1.5), Inches(3.6), Inches(2.8),
     "지도 시각화", [
         "",
         "영향권 오버레이",
         "민감지역 표시 (주거지, 학교)",
         "",
         "시나리오 A vs B:",
         '  "석탄 증가 → NOx +12%"',
         '  "LNG 기동 → NOx +3%"',
     ], accent=GREEN_PRIMARY, bg_color=GREEN_PALE)

# 가치
add_rounded(slide, Inches(0.8), Inches(4.7), Inches(11.7), Inches(2.3), GREEN_DARK)
tb(slide, Inches(1.3), Inches(4.85), Inches(10.7), Inches(0.35),
   "ESG 가치", size=18, color=MINT, bold=True)
tb(slide, Inches(1.3), Inches(5.25), Inches(10.7), Inches(1.5),
   "화력 가동을 늘릴 때의 환경 영향을 사전에 정량화\n\n"
   "운영자가 \"석탄을 더 돌릴까, LNG를 돌릴까\" 판단할 때,\n"
   "비용뿐 아니라 \"영흥 주변 주거지역 NOx +12%  vs  분당 주변 +3%\"를 함께 보여줌\n\n"
   "→ ESG 보고서에 정량적 근거 제공  /  지자체·주민 소통 자료",
   size=14, color=RGBColor(0xCC, 0xE5, 0xD9))

# ============================================================
# SLIDE 12: MVP & 로드맵
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
section_header(slide, "11", "MVP 범위 & 로드맵")

card(slide, Inches(0.8), Inches(1.5), Inches(5.5), Inches(3.3),
     "✅  MVP (Phase 1)", [
         "",
         "1. 태양광 하루전 예측 (ML 모델)",
         "2. 화력 출력 여유 분석",
         "3. 태양광 + 화력 겹쳐보기 시각화",
         "4. UC 기반 대응 계획 (액션 플랜)",
         "5. 시나리오 비교 (낙관/중립/비관)",
         "6. 가용용량 신고안 출력",
         "7. 기본 운영 대시보드",
         "8. 대기질 영향 지도 (기본)",
     ], accent=GREEN_PRIMARY, bg_color=GREEN_PALE)

card(slide, Inches(6.8), Inches(1.5), Inches(5.5), Inches(3.3),
     "🔜  Phase 2 (이후)", [
         "",
         "1. 실시간 감시 & 대응 모드",
         "2. 위성 기반 초단기 예측",
         "3. 베이지안+GRU 고급 모델",
         "4. 풍력 정밀 예측",
         "5. ESS 충방전 최적화 연계",
         "",
         "제외 항목:",
         "  KPX 자동 제출 / 송전망 제약",
         "  발전기 자동 제어 (사람이 판단)",
     ], accent=BLUE_INFO, bg_color=RGBColor(0xEA, 0xF4, 0xFB))

# 성공 지표
add_rounded(slide, Inches(0.8), Inches(5.2), Inches(11.7), Inches(1.8), LIGHT_BG)
tb(slide, Inches(1.2), Inches(5.3), Inches(10), Inches(0.35),
   "핵심 성공 지표", size=18, color=GREEN_DARK, bold=True)

metrics = [
    ("예측 정확도", "nRMSE 10% 이내", GREEN_PRIMARY),
    ("급변 탐지율", "80% 이상", BLUE_INFO),
    ("대응 실패", "0건", RED_ALERT),
    ("연료비 절감", "불필요 가동 감소", ORANGE_WARN),
]
for i, (title, value, color) in enumerate(metrics):
    x = Inches(1.2 + i * 2.8)
    add_rounded(slide, x, Inches(5.75), Inches(2.5), Inches(1.0), WHITE)
    add_rect(slide, x, Inches(5.75), Inches(2.5), Inches(0.05), color)
    tb(slide, x, Inches(5.85), Inches(2.5), Inches(0.3),
       title, size=14, color=color, bold=True, align=PP_ALIGN.CENTER)
    tb(slide, x, Inches(6.2), Inches(2.5), Inches(0.4),
       value, size=16, color=DARK_TEXT, bold=True, align=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 13: 마무리
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)

# 좌측 초록 패널
add_rect(slide, Inches(0), Inches(0), Inches(5.5), Inches(7.5), GREEN_DARK)
add_circle(slide, Inches(-0.5), Inches(5.5), Inches(3), RGBColor(0x0D, 0x4A, 0x38))

tb(slide, Inches(0.8), Inches(2.2), Inches(4), Inches(0.5),
   "One-line Summary", size=18, color=MINT)
tb(slide, Inches(0.8), Inches(2.8), Inches(4.2), Inches(2.5),
   "태양광 변동성을\n정밀 예측하고,\n화력에게 미리\n정확한 지시를 내린다.",
   size=32, color=WHITE, bold=True)

# 우측
tb(slide, Inches(6.5), Inches(2.0), Inches(6), Inches(0.5),
   "플랫폼이 답하는 질문", size=20, color=GREEN_DARK, bold=True)

questions = [
    "내일 태양광이 얼마나 나올까?",
    "급변하면 어떤 화력을 언제 올릴까?",
    "화력을 줄일 수 있는 시점은 언제인가?",
    "각 선택이 환경에 어떤 영향을 주는가?",
    "지금 예측이 빗나가고 있진 않은가?",
]
y = Inches(2.8)
for q in questions:
    add_circle(slide, Inches(6.5), y + Inches(0.08), Inches(0.15), GREEN_ACCENT)
    tb(slide, Inches(6.9), y, Inches(5.5), Inches(0.4), q, size=16, color=DARK_TEXT)
    y += Inches(0.55)

# 하단 버전
add_rect(slide, Inches(6.5), Inches(6.2), Inches(3), Inches(0.5), GREEN_PALE)
tb(slide, Inches(6.5), Inches(6.25), Inches(3), Inches(0.4),
   "  PRD v4  ·  2026.04", size=14, color=GREEN_PRIMARY, bold=True)


# 저장
output_path = "C:/Energy_effi/PRD_v4_presentation_v2.pptx"
prs.save(output_path)
print(f"OK: {output_path}")
print(f"Size: {os.path.getsize(output_path):,} bytes")
